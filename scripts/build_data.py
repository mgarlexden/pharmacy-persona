#!/usr/bin/env python3
"""reference/가상데이터_더미파일.xlsx -> data/floor12.json 변환 스크립트.

floor12.html(1~2층 개인 작업본)에 박혀 있던 DAYS·ITEMS 배열의 원본 데이터를
엑셀에서 직접 뽑아 JSON으로 분리한다. HTML에는 이 값을 다시 박지 않는다.

출처(floor12.html에는 주석이 없어 다시 확인한 것):
  날짜·요일·영업구분·연휴 전후·이비인후과 휴진·인플루엔자 지수·공휴일명 -> 00_공통달력
  내과/이비인후과 처방건수                                        -> L45
  처방 매출·비처방 매출·총매출                                    -> L46
  판콜에스(O003) 기말 실물재고·잔여일수                             -> L48 (품목코드 O003)
  조제장비 고장·직원 부재(그날 한 시간이라도 Y면 1)                  -> L55 (시간대별 → 일자별 OR)
  품목(진열 위치·마진 등)                                         -> L56 + L57 (품목코드로 합침, 페이스 수 > 0만)

L47(워크인 시간대)·L58·L59는 값을 직접 읽지 않는다. 시간대 가중치(RXW/WKW/SATW)와
개업 초기 계수 0.4는 floor12.html에 상수로 적힌 값을 그대로 쓴다(이 스크립트 밖, JS 쪽).

실행:  python scripts/build_data.py
"""
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "reference" / "가상데이터_더미파일.xlsx"
OUT = ROOT / "data" / "floor12.json"

# floor12.html의 SEASTXT(코드→문구)를 거꾸로 뒤집은 것. L56 '계절성' 열의 문구와
# 정확히 같은 글자여야 한다(다르면 KeyError로 바로 알 수 있게 둔다).
SEASON_KEY = {
    "겨울 감기 유행 연동": "cold",
    "감기 유행 약하게 연동": "coldw",
    "연휴 전날 1.5배": "eve15",
    "연휴 전날 3배": "eve3",
    "여름(7~8월) 1.5배": "sum78",
    "여름(6~8월) 1.3배": "sum68",
    "봄 황사·미세먼지 1.6배": "dust",
    "없음": "none",
}
WD = ["일", "월", "화", "수", "목", "금", "토"]
BIZ_CODE = {"휴무": 0, "평일": 1, "토요일": 2}
AFTER_CODE = {None: 0, "연휴 후 첫 영업일": 1, "연휴 전 마지막 영업일": 2}


def load_items(wb):
    l56 = {r[0]: r for r in wb["L56_상품구성·진열위치"].iter_rows(min_row=2, values_only=True) if r[0]}
    l57 = {r[0]: r for r in wb["L57_상품별마진"].iter_rows(min_row=2, values_only=True) if r[0]}

    items = []
    for code, r in l56.items():
        _, name, dept, category, zone, height, faces, weight, _pct, season_ko, _vendor = r
        if not faces or faces <= 0:
            continue  # 조제실 전용 등 손님이 보는 진열 품목이 아니다
        _, _, _, _, price, _cost, margin_pct, *_ = l57[code]
        items.append({
            "code": code,
            "name_ko": name,
            "department_ko": dept,
            "category_ko": category,
            "zone_ko": zone,
            "shelf_height_ko": height,
            "faces": faces,
            "sales_weight": weight,
            "season_key": SEASON_KEY[season_ko],
            "season_ko": season_ko,
            "price_krw": price,
            "margin_pct": margin_pct,
        })
    return items


def load_days(wb):
    cal = [r for r in wb["00_공통달력"].iter_rows(min_row=2, values_only=True) if r[0]]
    l45 = {r[0].date(): r for r in wb["L45_처방건수·조제료"].iter_rows(min_row=2, values_only=True) if r[0]}
    l46 = {r[0].date(): r for r in wb["L46_처방·일반약 매출비중"].iter_rows(min_row=2, values_only=True) if r[0]}
    l48_o3 = {
        r[0].date(): r for r in wb["L48_재고수량·잔여일수"].iter_rows(min_row=2, values_only=True)
        if r[0] and r[2] == "O003"
    }
    l55_flags = {}  # date -> (장비고장 있었는가, 직원부재 있었는가) — 하루 중 한 시간이라도 Y면 1
    for r in wb["L55_혼잡도·고객도착(시간별)"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        d = r[0].date()
        eq, stf = l55_flags.get(d, (False, False))
        l55_flags[d] = (eq or r[13] == "Y", stf or r[14] == "Y")

    days = []
    for row in cal:
        date, wd_ko, _mo, _iso, holiday_ko, _open_yn, biz_ko, _hours, after_ko, ent_closed_ko, ili, _resp = row
        d = date.date()
        rx = l45.get(d)
        rev = l46.get(d)
        inv = l48_o3.get(d)
        eq, stf = l55_flags.get(d, (False, False))
        days.append({
            "date": d.isoformat(),
            "weekday_ko": wd_ko,
            "weekday_idx": WD.index(wd_ko),
            "business_ko": biz_ko,
            "business_code": BIZ_CODE[biz_ko],
            "holiday_adjacent_ko": after_ko,
            "holiday_adjacent_code": AFTER_CODE[after_ko],
            "ent_closed": ent_closed_ko == "Y",
            "ili": ili,
            "im_rx_count": rx[3] if rx else 0,
            "ent_rx_count": rx[4] if rx else 0,
            "rx_revenue_krw": rev[3] if rev else 0,
            "otc_revenue_krw": rev[7] if rev else 0,
            "total_revenue_krw": rev[8] if rev else 0,
            "o003_closing_stock": inv[10] if inv else None,
            "o003_days_of_supply": inv[14] if inv else None,
            "holiday_name_ko": holiday_ko,
            "equipment_down": eq,
            "staff_absent": stf,
        })
    return days


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    items = load_items(wb)
    days = load_days(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps({"items": items, "days": days}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"  data/floor12.json  (items {len(items)}개, days {len(days)}일)")


if __name__ == "__main__":
    main()
